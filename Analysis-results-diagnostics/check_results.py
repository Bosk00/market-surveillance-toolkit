"""
Checks trade results from trades_log.txt against Polymarket.
Caches results in results_cache.json, outputs trades_analysis.xlsx
Run with: py -3.11 check_results.py
"""
import asyncio
import aiohttp
import json
import os
from collections import defaultdict
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CLOB_API  = "https://clob.polymarket.com"
GAMMA_API = "https://gamma-api.polymarket.com/markets/slug"
LOG_FILE   = "trades_log.txt"
CACHE_FILE = "results_cache.json"
OUT_FILE   = "trades_analysis.xlsx"

def load_cache():
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except: pass
    return {}

def save_cache(cache):
    with open(CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(cache, f, indent=2)

def parse_trades(filepath):
    trades = []
    try:
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
    except:
        print(f"Could not read {filepath}")
        return trades
    blocks = content.split("=" * 50)
    for block in blocks:
        if "Time:" not in block or "Token ID:" not in block:
            continue
        t = {}
        for line in block.strip().split("\n"):
            line = line.strip()
            if line.startswith("Time:"):
                t["time"] = line.replace("Time:", "").strip()
            elif line.startswith("Asset:"):
                t["asset"] = line.replace("Asset:", "").strip()
            elif line.startswith("Entry:"):
                try: t["entry"] = float(line.replace("Entry:", "").strip())
                except: pass
            # Accept both old ("Market ask:") and new ("Quoted ask:") label
            elif line.startswith("Market ask:") or line.startswith("Quoted ask:"):
                try: t["market_ask"] = float(line.split(":", 1)[1].strip())
                except: pass
            elif line.startswith("Slippage:"):
                try: t["slippage"] = float(line.split(":", 1)[1].strip())
                except: pass
            elif line.startswith("Size:"):
                t["size"] = line.replace("Size:", "").strip()
            elif line.startswith("Token ID:"):
                t["token_id"] = line.replace("Token ID:", "").strip()
            elif line.startswith("Confidence:"):
                t["conf"] = line.replace("Confidence:", "").strip()
        if "token_id" in t and "asset" in t:
            if len(t.get("time", "")) < 25 and len(t.get("token_id", "")) > 10:
                trades.append(t)
    return trades

async def check_token(session, token_id, trade_time_str, asset, direction):
    try:
        trade_dt  = datetime.strptime(trade_time_str, "%Y-%m-%d %H:%M:%S")
        trade_ts  = int(trade_dt.replace(tzinfo=timezone.utc).timestamp())
        window_ts = (trade_ts // 300) * 300
    except:
        return "ERROR", None
    asset_map = {
    "BTC":  "btc",
    "ETH":  "eth",
    "SOL":  "sol",
    "XRP":  "xrp",
    "DOGE": "doge",
    "BNB":  "bnb",
}
    asset_key = next((v for k, v in asset_map.items() if k in asset), None)
    if not asset_key: return "ERROR", None
    for offset in [14400, 14100, 14700, 18000, 10800, 0, 300, -300]:
        slug = f"{asset_key}-updown-5m-{window_ts + offset}"
        try:
            async with session.get(f"{GAMMA_API}/{slug}", timeout=aiohttp.ClientTimeout(total=5)) as resp:
                if resp.status != 200: continue
                data = await resp.json()
                outcome_prices = data.get("outcomePrices")
                closed   = data.get("closed", False)
                resolved = data.get("resolved", False)
                if not (closed or resolved):
                    try:
                        prices    = json.loads(outcome_prices) if isinstance(outcome_prices, str) else outcome_prices
                        our_price = float(prices[0] if "Up" in direction else prices[1])
                        if our_price >= 0.85:   return "WINNING", our_price
                        elif our_price <= 0.15: return "LOSING",  our_price
                        else:                   return "OPEN",    our_price
                    except: return "OPEN", None
                try:
                    prices    = json.loads(outcome_prices) if isinstance(outcome_prices, str) else outcome_prices
                    our_final = float(prices[0] if "Up" in direction else prices[1])
                    return ("WIN" if our_final >= 0.95 else "LOSS"), our_final
                except: return "UNCLEAR", None
        except: continue
    try:
        async with session.get(f"{CLOB_API}/book", params={"token_id": token_id}, timeout=aiohttp.ClientTimeout(total=5)) as resp:
            if resp.status == 200:
                data = await resp.json()
                asks = data.get("asks", []); bids = data.get("bids", [])
                best_ask = float(asks[0]["price"]) if asks else None
                best_bid = float(bids[0]["price"]) if bids else None
                if best_ask and best_ask >= 0.99: return "WIN",  best_bid
                elif best_bid and best_bid <= 0.02: return "LOSS", best_bid
    except: pass
    return "ERROR", None

def style_header(cell, bg="1F4E79"):
    cell.font      = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    cell.fill      = PatternFill("solid", start_color=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center")

def thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)

def get_conf_short(conf):
    """
    Classify confidence string into a display tier.
    REPRICED is checked first — repriced trades get their own tier
    regardless of the underlying signal type.
    """
    if "(REPRICED)" in conf: return "REPRICED"
    if "HIGH"   in conf:     return "HIGH"
    if "MEDIUM" in conf:     return "MED"
    if "LOW"    in conf:     return "LOW"
    if "PRICE"  in conf:     return "PRICE"
    return "?"

def build_excel(trade_results):
    wb = Workbook()

    # ── Sheet 1: All Trades ───────────────────────────────────────────
    ws = wb.active
    ws.title = "All Trades"

    # Col:  1      2       3        4            5             6        7            8           9          10        11      12        13
    headers = ["Date", "Time", "Asset", "Direction", "Confidence", "Entry", "Market Ask", "Slippage", "Size ($)", "Result", "P&L ($)", "Shares", "ROI (%)"]
    for col, h in enumerate(headers, 1):
        style_header(ws.cell(row=1, column=col, value=h))

    row = 2
    for t in trade_results:
        result    = t.get("result", "ERROR")
        entry     = t.get("entry", 0) or 0
        mkt_ask   = t.get("market_ask") or None
        slippage  = t.get("slippage")   or None
        # Derive slippage if we have both values but it wasn't stored directly
        if slippage is None and mkt_ask and entry and mkt_ask != entry:
            slippage = round(entry - mkt_ask, 4)
        size_str  = t.get("size", "$0")
        conf      = t.get("conf", "")
        time_str  = t.get("time", "")
        asset     = t.get("asset", "")

        conf_short  = get_conf_short(conf)
        direction   = "Up" if "Up" in asset else "Down"
        asset_name  = asset.replace(" Up", "").replace(" Down", "")

        try:    size_val = float(size_str.replace("$", ""))
        except: size_val = 0

        pnl = shares = roi = None
        if result == "WIN" and entry > 0 and size_val > 0:
            shares = round(size_val / entry, 2)
            pnl    = round(shares - size_val, 2)
            roi    = round(pnl / size_val * 100, 1)
        elif result == "LOSS" and size_val > 0:
            pnl = -size_val
            roi = -100.0

        date_part = time_str[:10] if time_str else ""
        time_part = time_str[11:19] if len(time_str) > 10 else ""

        vals = [
            date_part, time_part, asset_name, direction, conf_short,
            entry,
            round(mkt_ask, 4) if mkt_ask else None,
            round(slippage, 4) if slippage is not None else None,
            size_val, result, pnl, shares, roi,
        ]

        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font      = Font(name="Arial", size=9)
            c.border    = thin_border()
            c.alignment = Alignment(horizontal="center")

            # Result column (10)
            if col == 10:
                if result == "WIN":
                    c.fill = PatternFill("solid", start_color="C6EFCE")
                elif result == "LOSS":
                    c.fill = PatternFill("solid", start_color="FFC7CE")
                elif result in ("WINNING", "OPEN"):
                    c.fill = PatternFill("solid", start_color="FFEB9C")

            # P&L column (11)
            if col == 11 and pnl is not None:
                c.number_format = '#,##0.00'
                c.font = Font(name="Arial", size=9,
                              color="00B050" if pnl > 0 else "FF0000" if pnl < 0 else "000000")

            # Slippage column (8) — highlight bad slippage in amber
            if col == 8 and slippage is not None:
                c.number_format = '+0.0000;-0.0000;0'
                if slippage > 0.08:
                    c.fill = PatternFill("solid", start_color="FFC7CE")
                elif slippage > 0.03:
                    c.fill = PatternFill("solid", start_color="FFEB9C")

            # Confidence column (5) — colour REPRICED distinctly
            if col == 5 and conf_short == "REPRICED":
                c.fill = PatternFill("solid", start_color="E2EFDA")

        row += 1

    col_widths = [12, 10, 8, 10, 11, 8, 12, 10, 10, 10, 10, 8, 10]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{row - 1}"

    # ── Sheet 2: Summary by Date ──────────────────────────────────────
    ws2 = wb.create_sheet("By Date")
    date_headers = ["Date", "Trades", "Wins", "Losses", "Win Rate (%)", "Total P&L ($)", "Spent ($)", "ROI (%)"]
    for col, h in enumerate(date_headers, 1):
        style_header(ws2.cell(row=1, column=col, value=h))

    by_date = defaultdict(lambda: {"wins": 0, "losses": 0, "pnl": 0.0, "spent": 0.0, "total": 0})
    for t in trade_results:
        date  = t.get("time", "")[:10]
        res   = t.get("result", "")
        entry = t.get("entry", 0) or 0
        try: size_val = float(t.get("size", "$0").replace("$", ""))
        except: size_val = 0
        by_date[date]["total"] += 1
        by_date[date]["spent"] += size_val
        if res == "WIN" and entry > 0 and size_val > 0:
            by_date[date]["wins"] += 1
            by_date[date]["pnl"]  += round(size_val / entry - size_val, 2)
        elif res == "LOSS" and size_val > 0:
            by_date[date]["losses"] += 1
            by_date[date]["pnl"]    -= size_val

    for r, (date, s) in enumerate(sorted(by_date.items()), 2):
        total = s["wins"] + s["losses"]
        rate  = round(s["wins"] / total * 100, 1) if total > 0 else 0
        roi   = round(s["pnl"] / s["spent"] * 100, 1) if s["spent"] > 0 else 0
        vals  = [date, s["total"], s["wins"], s["losses"], rate, round(s["pnl"], 2), round(s["spent"], 2), roi]
        for col, v in enumerate(vals, 1):
            c = ws2.cell(row=r, column=col, value=v)
            c.font = Font(name="Arial", size=9)
            c.border = thin_border()
            c.alignment = Alignment(horizontal="center")
            if col == 6:
                c.number_format = '#,##0.00'
                if v > 0:  c.fill = PatternFill("solid", start_color="C6EFCE")
                elif v < 0: c.fill = PatternFill("solid", start_color="FFC7CE")
    for i, w in enumerate([12, 8, 6, 8, 12, 12, 10, 8], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    # ── Sheet 3: Summary by Confidence ───────────────────────────────
    ws3 = wb.create_sheet("By Confidence")
    conf_headers = ["Tier", "Trades", "Wins", "Losses", "Win Rate (%)", "Total P&L ($)", "Spent ($)", "ROI (%)", "Avg Slippage"]
    for col, h in enumerate(conf_headers, 1):
        style_header(ws3.cell(row=1, column=col, value=h))

    by_conf = defaultdict(lambda: {"wins": 0, "losses": 0, "pnl": 0.0, "spent": 0.0, "total": 0, "slippages": []})
    for t in trade_results:
        conf = t.get("conf", "")
        ck   = get_conf_short(conf)
        res  = t.get("result", "")
        entry = t.get("entry", 0) or 0
        slip  = t.get("slippage")
        mkt   = t.get("market_ask")
        if slip is None and mkt and entry and mkt != entry:
            slip = round(entry - mkt, 4)
        try: size_val = float(t.get("size", "$0").replace("$", ""))
        except: size_val = 0
        by_conf[ck]["total"] += 1
        by_conf[ck]["spent"] += size_val
        if slip is not None:
            by_conf[ck]["slippages"].append(slip)
        if res == "WIN" and entry > 0 and size_val > 0:
            by_conf[ck]["wins"] += 1
            by_conf[ck]["pnl"]  += round(size_val / entry - size_val, 2)
        elif res == "LOSS" and size_val > 0:
            by_conf[ck]["losses"] += 1
            by_conf[ck]["pnl"]    -= size_val

    # REPRICED first so it's easy to spot, then rest in priority order
    for r, ck in enumerate(["REPRICED", "HIGH", "MED", "PRICE", "LOW", "?"], 2):
        s = by_conf.get(ck)
        if not s: continue
        total    = s["wins"] + s["losses"]
        rate     = round(s["wins"] / total * 100, 1) if total > 0 else 0
        roi      = round(s["pnl"] / s["spent"] * 100, 1) if s["spent"] > 0 else 0
        avg_slip = round(sum(s["slippages"]) / len(s["slippages"]), 4) if s["slippages"] else None
        vals = [ck, s["total"], s["wins"], s["losses"], rate, round(s["pnl"], 2), round(s["spent"], 2), roi, avg_slip]
        for col, v in enumerate(vals, 1):
            c = ws3.cell(row=r, column=col, value=v)
            c.font = Font(name="Arial", size=9)
            c.border = thin_border()
            c.alignment = Alignment(horizontal="center")
            if col == 1 and ck == "REPRICED":
                c.fill = PatternFill("solid", start_color="E2EFDA")
                c.font = Font(name="Arial", size=9, bold=True)
            if col == 6:
                c.number_format = '#,##0.00'
                if v and v > 0:  c.fill = PatternFill("solid", start_color="C6EFCE")
                elif v and v < 0: c.fill = PatternFill("solid", start_color="FFC7CE")
    for i, w in enumerate([10, 8, 6, 8, 12, 12, 10, 8, 12], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # ── Sheet 4: By Asset ────────────────────────────────────────────
    ws4 = wb.create_sheet("By Asset")
    asset_headers = ["Asset", "Direction", "Trades", "Wins", "Losses", "Win Rate (%)", "P&L ($)", "Avg Entry", "Avg Slippage", "ROI (%)"]
    for col, h in enumerate(asset_headers, 1):
        style_header(ws4.cell(row=1, column=col, value=h))

    by_asset = defaultdict(lambda: {"wins": 0, "losses": 0, "pnl": 0.0, "spent": 0.0, "entries": [], "slippages": []})
    for t in trade_results:
        asset     = t.get("asset", "")
        direction = "Up" if "Up" in asset else "Down"
        aname     = asset.replace(" Up", "").replace(" Down", "")
        key       = (aname, direction)
        res       = t.get("result", "")
        entry     = t.get("entry", 0) or 0
        slip      = t.get("slippage")
        mkt       = t.get("market_ask")
        if slip is None and mkt and entry and mkt != entry:
            slip = round(entry - mkt, 4)
        try: size_val = float(t.get("size", "$0").replace("$", ""))
        except: size_val = 0
        by_asset[key]["spent"] += size_val
        if entry > 0:
            by_asset[key]["entries"].append(entry)
        if slip is not None:
            by_asset[key]["slippages"].append(slip)
        if res == "WIN" and entry > 0 and size_val > 0:
            by_asset[key]["wins"] += 1
            by_asset[key]["pnl"]  += round(size_val / entry - size_val, 2)
        elif res == "LOSS" and size_val > 0:
            by_asset[key]["losses"] += 1
            by_asset[key]["pnl"]    -= size_val

    for r, ((aname, direction), s) in enumerate(sorted(by_asset.items()), 2):
        total     = s["wins"] + s["losses"]
        rate      = round(s["wins"] / total * 100, 1) if total > 0 else 0
        roi       = round(s["pnl"] / s["spent"] * 100, 1) if s["spent"] > 0 else 0
        avg_entry = round(sum(s["entries"]) / len(s["entries"]), 3) if s["entries"] else 0
        avg_slip  = round(sum(s["slippages"]) / len(s["slippages"]), 4) if s["slippages"] else None
        vals = [aname, direction, total, s["wins"], s["losses"], rate, round(s["pnl"], 2), avg_entry, avg_slip, roi]
        for col, v in enumerate(vals, 1):
            c = ws4.cell(row=r, column=col, value=v)
            c.font = Font(name="Arial", size=9)
            c.border = thin_border()
            c.alignment = Alignment(horizontal="center")
            if col == 7:
                c.number_format = '#,##0.00'
                if v > 0:  c.fill = PatternFill("solid", start_color="C6EFCE")
                elif v < 0: c.fill = PatternFill("solid", start_color="FFC7CE")
    for i, w in enumerate([8, 10, 8, 6, 8, 12, 10, 10, 12, 8], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w

    # ── Sheet 5: Slippage Analysis ───────────────────────────────────
    ws5 = wb.create_sheet("Slippage")
    slip_headers = ["Date", "Time", "Asset", "Direction", "Confidence", "Market Ask", "Actual Entry", "Slippage", "Result", "P&L ($)"]
    for col, h in enumerate(slip_headers, 1):
        style_header(ws5.cell(row=1, column=col, value=h))

    slip_trades = []
    for t in trade_results:
        mkt  = t.get("market_ask")
        slip = t.get("slippage")
        entry = t.get("entry", 0) or 0
        if slip is None and mkt and entry and mkt != entry:
            slip = round(entry - mkt, 4)
        if slip is not None and abs(slip) > 0.001:
            slip_trades.append({**t, "_slip": slip})

    slip_trades.sort(key=lambda x: x.get("_slip", 0), reverse=True)

    for r, t in enumerate(slip_trades, 2):
        result   = t.get("result", "ERROR")
        entry    = t.get("entry", 0) or 0
        mkt_ask  = t.get("market_ask") or 0
        slip     = t.get("_slip", 0)
        size_str = t.get("size", "$0")
        conf     = t.get("conf", "")
        time_str = t.get("time", "")
        asset    = t.get("asset", "")

        conf_short  = get_conf_short(conf)
        direction   = "Up" if "Up" in asset else "Down"
        asset_name  = asset.replace(" Up", "").replace(" Down", "")

        try: size_val = float(size_str.replace("$", ""))
        except: size_val = 0

        pnl = None
        if result == "WIN" and entry > 0 and size_val > 0:
            pnl = round(size_val / entry - size_val, 2)
        elif result == "LOSS" and size_val > 0:
            pnl = -size_val

        date_part = time_str[:10] if time_str else ""
        time_part = time_str[11:19] if len(time_str) > 10 else ""

        vals = [date_part, time_part, asset_name, direction, conf_short,
                round(mkt_ask, 4) if mkt_ask else None, entry, round(slip, 4), result, pnl]

        for col, v in enumerate(vals, 1):
            c = ws5.cell(row=r, column=col, value=v)
            c.font = Font(name="Arial", size=9)
            c.border = thin_border()
            c.alignment = Alignment(horizontal="center")
            if col == 8:  # Slippage
                c.number_format = '+0.0000;-0.0000;0'
                if slip > 0.08:
                    c.fill = PatternFill("solid", start_color="FFC7CE")
                elif slip > 0.03:
                    c.fill = PatternFill("solid", start_color="FFEB9C")
            if col == 9:  # Result
                if result == "WIN":   c.fill = PatternFill("solid", start_color="C6EFCE")
                elif result == "LOSS": c.fill = PatternFill("solid", start_color="FFC7CE")
            if col == 10 and pnl is not None:
                c.number_format = '#,##0.00'
                c.font = Font(name="Arial", size=9,
                              color="00B050" if pnl > 0 else "FF0000" if pnl < 0 else "000000")

    for i, w in enumerate([12, 10, 8, 10, 11, 12, 12, 10, 10, 10], 1):
        ws5.column_dimensions[get_column_letter(i)].width = w
    ws5.freeze_panes = "A2"
    if len(slip_trades) > 0:
        ws5.auto_filter.ref = f"A1:J{len(slip_trades) + 1}"

    wb.save(OUT_FILE)
    print(f"Saved: {OUT_FILE}")

async def main():
    trades = parse_trades(LOG_FILE)
    print(f"Found {len(trades)} trades in {LOG_FILE}")
    if not trades:
        print("No trades found.")
        return

    cache   = load_cache()
    pending = [
        t for t in trades
        if t["token_id"] not in cache
        or cache[t["token_id"]]["result"] in ("OPEN", "WINNING", "LOSING", "ERROR")
    ]
    print(f"Cached: {len(cache)} | Need checking: {len(pending)}")

    if pending:
        print("Checking new/unresolved trades...")
        async with aiohttp.ClientSession() as session:
            for i, t in enumerate(pending):
                direction = "Up" if "Up" in t.get("asset", "") else "Down"
                result, _ = await check_token(
                    session, t["token_id"], t.get("time", ""), t.get("asset", ""), direction
                )
                # Derive slippage if not already stored
                entry   = t.get("entry", 0) or 0
                mkt_ask = t.get("market_ask", 0) or 0
                slippage = t.get("slippage")
                if slippage is None and mkt_ask and entry and mkt_ask != entry:
                    slippage = round(entry - mkt_ask, 4)

                cache[t["token_id"]] = {
                    "result":     result,
                    "time":       t.get("time", ""),
                    "asset":      t.get("asset", ""),
                    "entry":      entry,
                    "size":       t.get("size", "$0"),
                    "conf":       t.get("conf", ""),
                    "market_ask": mkt_ask,
                    "slippage":   slippage,
                }
                if i % 20 == 19:
                    await asyncio.sleep(0.5)
                    save_cache(cache)
        save_cache(cache)
        print(f"Cache updated: {len(cache)} entries")

    log_by_token = {t["token_id"]: t for t in trades}

    trade_results = []
    seen = set()
    for tid, cached in cache.items():
        if tid in seen:
            continue
        seen.add(tid)
        if tid in log_by_token:
            trade_results.append({**log_by_token[tid], **cached})
        else:
            trade_results.append(cached)

    trade_results.sort(key=lambda x: x.get("time", ""))
    build_excel(trade_results)

if __name__ == "__main__":
    asyncio.run(main())